"""
kml_to_tabla.py
================
Utilidades genéricas para convertir archivos KML/KMZ a tablas (pandas DataFrames)
y exportarlas a Excel, una hoja por carpeta/capa del KML.

Pensado para ejecutarse en un notebook (Jupyter / Colab / VSCode).

Instalación de dependencias (si hace falta):
    pip install lxml pandas openpyxl

Uso rápido:
    from kml_to_tabla import kml_or_kmz_to_dataframes, dataframes_to_excel

    capas = kml_or_kmz_to_dataframes("diagrama_con_pozos_ypf.kml")
    dataframes_to_excel(capas, "ypf_tabla.xlsx")

    capas_pae = kml_or_kmz_to_dataframes("pozos_pae.kmz")
    dataframes_to_excel(capas_pae, "pae_tabla.xlsx")
"""

import os
import zipfile
import tempfile
import html
import pandas as pd
from lxml import etree
import lxml.html as LH
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl import Workbook

# Namespace estándar de KML. Algunos archivos usan otros namespaces/gx, pero
# la etiqueta base kml/2.2 es casi universal.
NS = {"kml": "http://www.opengis.net/kml/2.2"}


# ---------------------------------------------------------------------------
# 1. Carga del archivo (soporta .kml directo y .kmz comprimido)
# ---------------------------------------------------------------------------

def _load_kml_bytes(path):
    """Devuelve los bytes del .kml, sea que 'path' apunte a un .kml o a un .kmz."""
    if path.lower().endswith(".kmz"):
        with zipfile.ZipFile(path, "r") as z:
            # El KML principal suele llamarse doc.kml, pero por las dudas
            # tomamos el primer .kml que aparezca dentro del zip.
            kml_names = [n for n in z.namelist() if n.lower().endswith(".kml")]
            if not kml_names:
                raise ValueError(f"No se encontró ningún .kml dentro de {path}")
            with z.open(kml_names[0]) as f:
                return f.read()
    else:
        with open(path, "rb") as f:
            return f.read()


def _parse_root(kml_bytes):
    """
    Parsea el XML con recover=True porque muchos KML "reales" (exportados desde
    Google Earth / ArcGIS) tienen namespaces mal declarados (ej: prefijo xsi sin
    declarar) que rompen un parser XML estricto.
    """
    parser = etree.XMLParser(recover=True)
    root = etree.fromstring(kml_bytes, parser=parser)
    return root


# ---------------------------------------------------------------------------
# 2. Extracción de atributos de cada Placemark (con fallbacks)
# ---------------------------------------------------------------------------

def _parse_extended_data(placemark):
    """
    Caso 1: atributos estructurados en <ExtendedData>.
    Soporta tanto <Data name="X"><value>Y</value></Data>
    como <SimpleData name="X">Y</SimpleData> (usado junto con <Schema>).
    """
    d = {}
    ext = placemark.find("kml:ExtendedData", NS)
    if ext is None:
        return d

    for data in ext.findall("kml:Data", NS):
        key = data.get("name")
        val_el = data.find("kml:value", NS)
        val = val_el.text if val_el is not None else None
        if key:
            d[key] = (val or "").strip()

    for simple in ext.findall(".//kml:SimpleData", NS):
        key = simple.get("name")
        val = simple.text
        if key:
            d[key] = (val or "").strip()

    return d


def _parse_description_table(desc_text):
    """
    Caso 2: atributos como tabla HTML dentro de <description> (típico de
    exports de ArcGIS/Esri). Usamos un parser HTML real (no regex) porque
    estas tablas suelen venir anidadas (una tabla "contenedora" con el título
    y adentro la tabla real de pares clave-valor), y un regex ingenuo mezcla
    ambos niveles.
    """
    d = {}
    if not desc_text:
        return d
    try:
        doc = LH.fromstring(desc_text)
    except Exception:
        return d

    for tr in doc.findall(".//tr"):
        tds = tr.findall("td")
        if len(tds) == 2:
            key = (tds[0].text_content() or "").strip()
            val = (tds[1].text_content() or "").strip()
            if key:
                d[key] = val
    return d


def _parse_plain_description(placemark):
    """
    Caso 3 (fallback final): si <description> no es HTML tabular, la
    guardamos como texto plano en una sola columna para no perder información.
    """
    desc = placemark.find("kml:description", NS)
    if desc is None or not desc.text:
        return {}
    text = html.unescape(desc.text).strip()
    # Si parece HTML (contiene tags) y ya lo intentamos parsear como tabla
    # y no salió nada, no lo repetimos como texto crudo lleno de tags.
    if "<" in text and ">" in text:
        return {}
    return {"Descripcion": text}


def _get_geometry(placemark):
    """
    Devuelve (tipo_geometria, lon, lat, elevacion, info_extra).
    Soporta Point, LineString y Polygon (incluye MultiGeometry porque
    buscamos con './/').
    """
    pt = placemark.find(".//kml:Point/kml:coordinates", NS)
    if pt is not None and pt.text:
        parts = pt.text.strip().split(",")
        lon = parts[0] if len(parts) > 0 else ""
        lat = parts[1] if len(parts) > 1 else ""
        elev = parts[2] if len(parts) > 2 else ""
        return "Point", lon, lat, elev, None

    ls = placemark.find(".//kml:LineString/kml:coordinates", NS)
    if ls is not None and ls.text:
        n_vertices = len(ls.text.strip().split())
        return "LineString", "", "", "", f"{n_vertices} vértices"

    poly = placemark.find(".//kml:Polygon//kml:coordinates", NS)
    if poly is not None and poly.text:
        n_vertices = len(poly.text.strip().split())
        return "Polygon", "", "", "", f"{n_vertices} vértices"

    return "Sin geometría", "", "", "", None


def _placemark_to_row(placemark):
    """Combina nombre + atributos (con fallbacks) + geometría en un dict (fila)."""
    name_el = placemark.find("kml:name", NS)
    name = name_el.text.strip() if (name_el is not None and name_el.text) else ""

    # Probamos las 3 estrategias en orden; si ExtendedData no trae nada,
    # probamos la tabla HTML; si tampoco, guardamos el texto plano.
    attrs = _parse_extended_data(placemark)
    if not attrs:
        desc = placemark.find("kml:description", NS)
        attrs = _parse_description_table(desc.text if desc is not None else "")
    if not attrs:
        attrs = _parse_plain_description(placemark)

    geom_type, lon, lat, elev, extra = _get_geometry(placemark)

    row = {"Nombre_KML": name}
    row.update(attrs)
    row["Geometria"] = geom_type
    if geom_type == "Point":
        row["Longitud"] = lon
        row["Latitud"] = lat
        row["Elevacion"] = elev
    elif extra:
        row["Info_geometria"] = extra

    return row


# ---------------------------------------------------------------------------
# 3. Recorrido de carpetas -> un DataFrame por carpeta
# ---------------------------------------------------------------------------

def kml_or_kmz_to_dataframes(path, agrupar_por_carpeta=True):
    """
    Convierte un archivo .kml o .kmz en un diccionario {nombre_capa: DataFrame}.

    Parámetros
    ----------
    path : str
        Ruta al archivo .kml o .kmz
    agrupar_por_carpeta : bool
        Si True (default), genera un DataFrame por cada <Folder> del KML
        (esto es lo más útil cuando el archivo tiene varias capas: pozos,
        caminos, instalaciones, etc.).
        Si False, devuelve un único DataFrame con todos los Placemarks del
        archivo (útil para KML simples, de una sola capa, como el de PAE).

    Retorna
    -------
    dict[str, pandas.DataFrame]
    """
    kml_bytes = _load_kml_bytes(path)
    root = _parse_root(kml_bytes)

    resultado = {}

    if agrupar_por_carpeta:
        folders = root.findall(".//kml:Folder", NS)
        if not folders:
            # No hay carpetas: tratamos todo el documento como una sola capa.
            agrupar_por_carpeta = False
        else:
            for folder in folders:
                name_el = folder.find("kml:name", NS)
                folder_name = name_el.text.strip() if (name_el is not None and name_el.text) else "Capa_sin_nombre"
                placemarks = folder.findall("kml:Placemark", NS)
                if not placemarks:
                    continue
                rows = [_placemark_to_row(pm) for pm in placemarks]
                resultado[folder_name] = pd.DataFrame(rows)

    if not agrupar_por_carpeta:
        placemarks = root.findall(".//kml:Placemark", NS)
        rows = [_placemark_to_row(pm) for pm in placemarks]
        nombre_doc_el = root.find(".//kml:Document/kml:name", NS)
        nombre_doc = nombre_doc_el.text.strip() if (nombre_doc_el is not None and nombre_doc_el.text) else "Datos"
        resultado[nombre_doc[:31]] = pd.DataFrame(rows)

    return resultado


# ---------------------------------------------------------------------------
# 4. Exportar a Excel con formato prolijo
# ---------------------------------------------------------------------------

def dataframes_to_excel(capas: dict, output_path: str):
    """
    Escribe un diccionario {nombre_hoja: DataFrame} a un archivo Excel,
    con encabezados formateados, filtros automáticos y columnas de
    lat/lon al frente.
    """
    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    BODY_FONT = Font(name="Arial", size=10)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, df in capas.items():
        if df.empty:
            continue

        ws = wb.create_sheet(str(sheet_name)[:31])

        # Columnas prioritarias primero (si existen)
        cols = list(df.columns)
        priority = [c for c in ["Nombre_KML", "Latitud", "Longitud", "Elevacion"] if c in cols]
        rest = [c for c in cols if c not in priority]
        ordered_cols = priority + rest
        df = df[ordered_cols]

        for j, col in enumerate(ordered_cols, start=1):
            c = ws.cell(row=1, column=j, value=col)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(vertical="center")

        for i, row in enumerate(df.itertuples(index=False), start=2):
            for j, val in enumerate(row, start=1):
                ws.cell(row=i, column=j, value=val).font = BODY_FONT

        for j, col in enumerate(ordered_cols, start=1):
            muestra = df[col].astype(str).values[:200]
            maxlen = max([len(str(col))] + [len(v) for v in muestra])
            ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 2, 10), 45)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(ordered_cols))}{len(df) + 1}"

    wb.save(output_path)
    print(f"Archivo guardado: {output_path}")


# ---------------------------------------------------------------------------
# 5. Ejemplo de uso (correr en un notebook)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    # --- YPF: archivo con muchas capas (pozos, locaciones, ductos, etc.) ---
    capas_ypf = kml_or_kmz_to_dataframes("diagrama_con_pozos_ypf.kml", agrupar_por_carpeta=True)
    for nombre, df in capas_ypf.items():
        print(f"{nombre}: {len(df)} filas, {len(df.columns)} columnas")
    dataframes_to_excel(capas_ypf, "ypf_tabla.xlsx")

    # --- PAE: archivo simple, una sola capa/placemark ---
    capas_pae = kml_or_kmz_to_dataframes("pozos_pae.kmz", agrupar_por_carpeta=False)
    for nombre, df in capas_pae.items():
        print(f"{nombre}: {len(df)} filas, {len(df.columns)} columnas")
    dataframes_to_excel(capas_pae, "pae_tabla.xlsx")